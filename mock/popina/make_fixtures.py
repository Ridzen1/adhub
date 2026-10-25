# Découpe les vrais exports (21-22/10/2026) en un fichier par jour pour le faux serveur.
import openpyxl, sys, os
D = os.path.dirname(os.path.abspath(__file__)) + '/fixtures'
def split(kind, header_rows, day_col):
    src = openpyxl.load_workbook(f'{D}/{kind}_20261021_20261022.xlsx').active
    rows = [list(r) for r in src.iter_rows(values_only=True)]
    head, body = rows[:header_rows], rows[header_rows:]
    for day in ['21/10/2026', '22/10/2026']:
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Basic Worksheet'
        for r in head: ws.append(r)
        for r in body:
            if isinstance(r[day_col], str) and r[day_col].startswith(day): ws.append(r)
        ws.append(['Total'])
        d, m, y = day.split('/')
        wb.save(f'{D}/{kind}_{y}-{m}-{d}.xlsx')
    # version vide (jour sans service)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Basic Worksheet'
    for r in head: ws.append(r)
    ws.append(['Total']); wb.save(f'{D}/{kind}_vide.xlsx')
split('accounting', 2, 0)
split('items', 1, 1)
print(sorted(os.listdir(D)))
